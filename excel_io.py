"""Excel reading (xlsx/xls) and writing results into a copy of the workbook."""
import datetime
import os
import re
from copy import copy

import openpyxl
from openpyxl.utils import get_column_letter

OUTPUT_COLUMNS = ["Geocoding_Status", "Matched_Address", "Location_Type", "Place_ID", "Geocoding_Note"]

FIELD_KEYWORDS = {
    "name": ["restaurant", "customer", "outlet", "party", "shop", "business", "name"],
    "address": ["address", "addr", "street", "location"],
    "city": ["city", "town"],
    "state": ["state"],
    "pin": ["pincode", "pin", "zip", "postal"],
    "country": ["country"],
    "lat": ["latitude", "lat"],
    "lng": ["longitude", "lng", "lon", "long"],
}


def convert_xls(src, dst):
    """Convert legacy .xls to .xlsx (values only; .xls formatting cannot be carried over)."""
    import xlrd
    book = xlrd.open_workbook(src)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sh in book.sheets():
        ws = wb.create_sheet(sh.name[:31])
        for r in range(sh.nrows):
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_EMPTY:
                    continue
                v = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        v = xlrd.xldate_as_datetime(v, book.datemode)
                    except Exception:
                        pass
                elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                    v = bool(v)
                elif cell.ctype == xlrd.XL_CELL_ERROR:
                    v = None
                ws.cell(row=r + 1, column=c + 1, value=v)
    wb.save(dst)


def sheet_names(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return wb.sheetnames
    finally:
        wb.close()


def _read_values(path, sheet):
    """All cell values of a sheet as a list of tuples (formulas -> last computed values)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return [tuple(r) for r in wb[sheet].iter_rows(values_only=True)]
    finally:
        wb.close()


def _filled(v):
    return v is not None and str(v).strip() != ""


def detect_header_row(rows):
    """First row (1-based) within the first 20 with at least 2 text cells."""
    for i, row in enumerate(rows[:20]):
        if sum(1 for v in row if isinstance(v, str) and v.strip()) >= 2:
            return i + 1
    return 1


def _display(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.isoformat()
    return str(v)


def guess_mapping(headers):
    """Best-effort column guess from header names (user can change it)."""
    mapping, used = {}, set()
    for field, words in FIELD_KEYWORDS.items():
        for word in words:
            hit = next((i for i, h in enumerate(headers)
                        if i not in used and re.search(r"(^|[^a-z])" + word + r"([^a-z]|$)", h.lower())), None)
            if hit is None:
                hit = next((i for i, h in enumerate(headers) if i not in used and word in h.lower()), None)
            if hit is not None:
                mapping[field] = hit
                used.add(hit)
                break
    return mapping


class SheetData:
    """Values of one worksheet plus header info."""

    def __init__(self, path, sheet, header_row=None):
        self.path, self.sheet = path, sheet
        rows = _read_values(path, sheet)
        self.header_row = header_row or detect_header_row(rows)
        hdr = rows[self.header_row - 1] if len(rows) >= self.header_row else ()
        width = max((len(r) for r in rows), default=0)
        last = max((i + 1 for i, v in enumerate(hdr) if _filled(v)), default=0)
        self.header_width = last  # columns that have a header
        self.headers = [(_display(hdr[i]).strip() if i < len(hdr) and _filled(hdr[i]) else "")
                        or f"(Column {get_column_letter(i + 1)})" for i in range(width)]
        # Data rows: (excel_row_number, values) excluding fully blank rows
        self.rows = []
        for i, r in enumerate(rows[self.header_row:], start=self.header_row + 1):
            if any(_filled(v) for v in r):
                self.rows.append((i, tuple(r) + (None,) * (width - len(r))))

    def preview(self, n=5):
        return {
            "headers": [f"{get_column_letter(i + 1)}: {h}" for i, h in enumerate(self.headers)],
            "sample": [[_display(v) for v in vals] for _, vals in self.rows[:n]],
            "header_row": self.header_row,
            "row_count": len(self.rows),
            "guess": guess_mapping(self.headers),
        }


def write_output(src_path, dst_path, sheet, header_row, results, lat_col=None, lng_col=None):
    """Copy the workbook and write results into the chosen sheet.

    results: {excel_row: {"lat", "lng", "status", "matched", "loc_type", "place_id", "note"}}
    lat_col/lng_col: 0-based existing columns to reuse; otherwise new columns are added.
    Only the result columns are written; every other cell is left untouched.
    """
    wb = openpyxl.load_workbook(src_path)
    ws = wb[sheet]
    hdr_cells = list(ws[header_row]) if ws.max_row >= header_row else []
    existing = {str(c.value).strip().lower(): c.column for c in hdr_cells if _filled(c.value)}
    last_col = max([c.column for c in hdr_cells if _filled(c.value)] or [0])
    style_src = ws.cell(row=header_row, column=last_col) if last_col else None

    def col_for(name, given):
        nonlocal last_col
        if given is not None:
            return given + 1
        if name.lower() in existing:
            return existing[name.lower()]
        last_col += 1
        c = ws.cell(row=header_row, column=last_col, value=name)
        if style_src is not None and style_src.has_style:
            c.font, c.fill, c.border, c.alignment = (copy(style_src.font), copy(style_src.fill),
                                                     copy(style_src.border), copy(style_src.alignment))
        ws.column_dimensions[get_column_letter(last_col)].width = 40 if name == "Matched_Address" else 18
        return last_col

    cols = {"lat": col_for("Latitude", lat_col), "lng": col_for("Longitude", lng_col)}
    for name in OUTPUT_COLUMNS:
        cols[name] = col_for(name, None)

    for row, res in results.items():
        if res.get("lat") is not None and res.get("lng") is not None:
            for key in ("lat", "lng"):
                c = ws.cell(row=row, column=cols[key], value=round(float(res[key]), 7))
                c.number_format = "0.0000000"
        # Failed lookups never touch existing coordinates (no fake/blank overwrite)
        ws.cell(row=row, column=cols["Geocoding_Status"], value=res.get("status") or None)
        ws.cell(row=row, column=cols["Matched_Address"], value=res.get("matched") or None)
        ws.cell(row=row, column=cols["Location_Type"], value=res.get("loc_type") or None)
        ws.cell(row=row, column=cols["Place_ID"], value=res.get("place_id") or None)
        ws.cell(row=row, column=cols["Geocoding_Note"], value=res.get("note") or None)

    os.makedirs(os.path.dirname(dst_path) or ".", exist_ok=True)
    wb.save(dst_path)
