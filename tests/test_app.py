import io
import os
import sys

import openpyxl
import pytest
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import app as appmod  # noqa: E402
import geocoder as g  # noqa: E402
import fake_google  # noqa: E402

HEAD = ["Cust ID", "Restaurant Name", "Phone", "Full Address", "City", "State", "Pincode"]
ROWS = [
    [101, "Cafe Mocha", "9876500001", "12 Koregaon Park Lane 5", "Pune", "Maharashtra", 411001],
    [102, "Hotel Shreyas", "9876500002", "Plot 4, Baner Road", "Pune", "Maharashtra", 411045],
    [103, "Lost Dhaba", "9876500003", "nowhere village xyz", "Pune", "Maharashtra", 411001],
]


def make_xlsx(path, sheets):
    """sheets: {name: (headers, rows)}; first sheet header gets bold + fill to check formatting."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for i, (name, (head, rows)) in enumerate(sheets.items()):
        ws = wb.create_sheet(name)
        ws.append(head)
        for r in rows:
            ws.append(r)
        if i == 0:
            for c in ws[1]:
                c.font = Font(bold=True)
                c.fill = PatternFill("solid", fgColor="FFFF00")
            ws.column_dimensions["B"].width = 33
    wb.save(path)
    return path


@pytest.fixture
def env(tmp_path, monkeypatch):
    monkeypatch.setattr(appmod, "WORK_DIR", str(tmp_path / "work"))
    monkeypatch.setattr(appmod, "CACHE_PATH", str(tmp_path / "work" / "cache.json"))
    monkeypatch.setitem(appmod.KEY, "value", "test-key")
    fake_google.calls.clear()
    holder = {}

    def fake_geo():
        if "geo" not in holder:
            os.makedirs(appmod.WORK_DIR, exist_ok=True)
            holder["geo"] = g.Geocoder(appmod.KEY["value"], appmod.CACHE_PATH, rps=0,
                                       session=fake_google.FakeSession(), sleep=lambda s: None)
        return holder["geo"]

    monkeypatch.setattr(appmod, "get_geocoder", fake_geo)
    appmod.S.reset()
    client = appmod.app.test_client()
    client.tmp = tmp_path
    client.holder = holder
    yield client
    if appmod.S.running():
        appmod.S.stop_event.set()
        appmod.S.thread.join()
    appmod.S.reset()


def upload(client, path):
    with open(path, "rb") as f:
        r = client.post("/api/upload", data={"file": (io.BytesIO(f.read()), os.path.basename(path))})
    assert r.status_code == 200, r.json
    return r.json


def preview(client, sheet):
    r = client.get(f"/api/preview?sheet={sheet}")
    assert r.status_code == 200, r.json
    return r.json


def run(client, body, mode="all"):
    r = client.post("/api/start", json={**body, "mode": mode})
    assert r.status_code == 200, r.json
    if appmod.S.thread:
        appmod.S.thread.join(timeout=120)
    return client.get("/api/status").json


def download(client, path):
    r = client.get("/api/download")
    assert r.status_code == 200, r.json
    with open(path, "wb") as f:
        f.write(r.data)
    assert "_geocoded.xlsx" in r.headers["Content-Disposition"]
    return openpyxl.load_workbook(path)


def body_for(sheet, mapping, **kw):
    return {"sheet": sheet, "mapping": mapping, "skip_existing": True, "default_country": "India", **kw}


STD_MAP = {"name": 1, "address": 3, "city": 4, "state": 5, "pin": 6}


def col(ws, name):
    return next(c.column for c in ws[1] if c.value == name)


# ---------- unit tests ----------

def test_build_address_dedup_and_queries():
    parts = {"name": "Cafe Mocha", "address": "12 MG Road, Pune", "city": "Pune", "state": "Maharashtra",
             "pin": 411001.0, "country": None}
    assert g.build_address(parts, "India") == "12 MG Road, Pune, Maharashtra, 411001, India"
    qs = g.build_queries(parts, "India")
    assert qs == ["12 MG Road, Pune, Maharashtra, 411001, India",
                  "Cafe Mocha, 12 MG Road, Pune, Maharashtra, 411001, India"]
    assert g.build_queries({"name": "X", "state": "Maharashtra"}, "India") == []  # not enough
    assert g.build_queries({"name": "X", "city": "Pune"}, "India") == ["X, Pune, India"]


def _r(lt, types, lat=18.5, lng=73.8, postal="411001", partial=False):
    return {"formatted_address": "a", "location": {"lat": lat, "lng": lng}, "location_type": lt, "types": types,
            "partial_match": partial, "place_id": "p", "postal_code": postal}


def test_classify():
    assert g.classify([_r("ROOFTOP", ["premise"])], "411001")[0] == g.PRECISE
    assert g.classify([_r("RANGE_INTERPOLATED", ["street_address"])])[0] == g.APPROXIMATE
    assert g.classify([_r("GEOMETRIC_CENTER", ["route"])])[0] == g.APPROXIMATE
    assert g.classify([_r("APPROXIMATE", ["locality", "political"])])[0] == g.REVIEW
    assert g.classify([_r("APPROXIMATE", ["postal_code"])])[0] == g.REVIEW
    assert g.classify([_r("ROOFTOP", ["premise"], partial=True)])[0] == g.APPROXIMATE
    assert g.classify([_r("ROOFTOP", ["premise"], postal="400001")], "411001")[0] == g.REVIEW
    assert g.classify([_r("ROOFTOP", ["premise"]), _r("ROOFTOP", ["premise"], lat=19.5)])[0] == g.REVIEW
    assert g.classify([_r("ROOFTOP", ["premise"]), _r("ROOFTOP", ["premise"], lat=18.5001)])[0] == g.PRECISE
    assert g.classify([])[0] == g.NOT_FOUND


def test_retry_on_unknown_error():
    geo = g.Geocoder("k", rps=0, session=fake_google.FakeSession(), sleep=lambda s: None)
    out = geo.lookup("flaky place 1, Pune")
    assert out["status"] == "OK" and geo.requests_made == 2


def test_invalid_request_is_api_error_not_fatal():
    geo = g.Geocoder("k", rps=0, session=fake_google.FakeSession(), sleep=lambda s: None)
    assert geo.lookup("invalid thing")["status"] == "INVALID_REQUEST"


def test_network_errors_retry_then_api_error():
    import requests

    class Boom:
        n = 0

        def get(self, *a, **k):
            Boom.n += 1
            raise requests.ConnectionError("down")

    geo = g.Geocoder("k", rps=0, session=Boom(), sleep=lambda s: None, max_retries=2)
    out = geo.lookup("12 Some Road, Pune")
    assert out["status"] == "API_ERROR" and Boom.n == 3


# ---------- end-to-end (test cases 1-10) ----------

def test_1_complete_addresses_and_preservation(env):
    src = make_xlsx(env.tmp / "customers.xlsx", {"Customers": (HEAD, ROWS), "Notes": (["a", "b"], [[1, 2]])})
    assert upload(env, src)["sheets"] == ["Customers", "Notes"]
    pv = preview(env, "Customers")
    assert pv["guess"] == {"name": 1, "address": 3, "city": 4, "state": 5, "pin": 6}
    st = run(env, body_for("Customers", pv["guess"]))
    assert st["state"] == "done"
    assert st["counts"] == {g.PRECISE: 2, g.NOT_FOUND: 1}

    wb = download(env, env.tmp / "out.xlsx")
    ws = wb["Customers"]
    # original data unchanged
    orig = openpyxl.load_workbook(src)["Customers"]
    for r in range(1, 5):
        for c in range(1, 8):
            assert ws.cell(r, c).value == orig.cell(r, c).value
    # formatting, widths, other sheet kept
    assert ws["A1"].font.bold and ws.column_dimensions["B"].width == 33
    assert wb.sheetnames == ["Customers", "Notes"] and wb["Notes"]["A2"].value == 1
    # coordinates belong to the right row
    lat, lng = col(ws, "Latitude"), col(ws, "Longitude")
    for i, row in enumerate(ROWS[:2], start=2):
        q = g.build_queries({"name": row[1], "address": row[3], "city": row[4], "state": row[5], "pin": row[6]}, "India")[0]
        exp = fake_google._coord(q.lower())
        assert ws.cell(i, lat).value == pytest.approx(exp[0], abs=1e-6)
        assert ws.cell(i, lng).value == pytest.approx(exp[1], abs=1e-6)
        assert ws.cell(i, col(ws, "Geocoding_Status")).value == g.PRECISE
        assert ws.cell(i, col(ws, "Matched_Address")).value.startswith("Matched:")
        assert ws.cell(i, col(ws, "Location_Type")).value == "ROOFTOP"
        assert ws.cell(i, col(ws, "Place_ID")).value.startswith("PID_")
    # not found: no fake coordinates
    assert ws.cell(4, lat).value is None and ws.cell(4, lng).value is None
    assert ws.cell(4, col(ws, "Geocoding_Status")).value == g.NOT_FOUND
    # original file never modified (upload is a copy)
    assert openpyxl.load_workbook(src)["Customers"].max_column == 7


def test_2_different_names_and_order(env):
    head = ["PIN", "Town", "Outlet", "Addr Line", "Mobile"]
    rows = [[411001, "Pune", "Bake House", "5 Camp Area Lane", "999"]]
    src = make_xlsx(env.tmp / "c2.xlsx", {"S": (head, rows)})
    upload(env, src)
    pv = preview(env, "S")
    assert pv["guess"]["pin"] == 0 and pv["guess"]["city"] == 1 and pv["guess"]["name"] == 2 and pv["guess"]["address"] == 3
    st = run(env, body_for("S", {"pin": 0, "city": 1, "name": 2, "address": 3}))
    assert st["counts"] == {g.PRECISE: 1}
    ws = download(env, env.tmp / "o.xlsx")["S"]
    assert ws.cell(2, col(ws, "Latitude")).value is not None and ws["E2"].value == "999"


def test_3_missing_addresses_and_blank_rows(env):
    rows = [ROWS[0], [None] * 7, [104, "No Addr Cafe", "1", None, None, "Maharashtra", None], ROWS[1]]
    src = make_xlsx(env.tmp / "c3.xlsx", {"S": (HEAD, rows)})
    upload(env, src)
    assert preview(env, "S")["row_count"] == 3  # blank row ignored
    st = run(env, body_for("S", STD_MAP))
    assert st["counts"] == {g.PRECISE: 2, g.MISSING: 1}
    ws = download(env, env.tmp / "o.xlsx")["S"]
    s = col(ws, "Geocoding_Status")
    assert ws.cell(3, s).value is None  # blank row untouched
    assert ws.cell(4, s).value == g.MISSING and ws.cell(4, col(ws, "Latitude")).value is None
    assert ws.cell(5, s).value == g.PRECISE


def test_4_duplicate_addresses_use_cache(env):
    dup = [105, "Other Cafe", "2", "12 Koregaon Park Lane 5", "Pune", "Maharashtra", 411001]
    src = make_xlsx(env.tmp / "c4.xlsx", {"S": (HEAD, [ROWS[0], dup])})
    upload(env, src)
    est = env.post("/api/estimate", json=body_for("S", STD_MAP)).json
    assert est["min_requests"] == 1 and est["duplicates"] == 2
    st = run(env, body_for("S", STD_MAP))
    assert len(fake_google.calls) == 1  # second row served from cache
    assert st["counts"] == {g.PRECISE: 2}
    notes = {r["row"]: r["note"] for r in st["results"]}
    assert "Same address as row(s) 3" in notes[2] and "Same address as row(s) 2" in notes[3]


def test_5_existing_coordinates(env):
    head = HEAD + ["Latitude", "Longitude"]
    rows = [ROWS[0] + [18.1, 73.1], ROWS[1] + [None, None], ROWS[2] + [18.2, 73.2]]
    src = make_xlsx(env.tmp / "c5.xlsx", {"S": (head, rows)})
    upload(env, src)
    pv = preview(env, "S")
    assert pv["guess"]["lat"] == 7 and pv["guess"]["lng"] == 8
    st = run(env, body_for("S", pv["guess"]))
    assert st["counts"] == {g.SKIPPED: 2, g.PRECISE: 1}
    ws = download(env, env.tmp / "o.xlsx")["S"]
    assert ws.max_column == 9 + 5  # reused Latitude/Longitude, no duplicates
    assert (ws["H2"].value, ws["I2"].value) == (18.1, 73.1) and ws["H3"].value is not None

    # Skip disabled: re-geocode; a failed lookup must keep previous coordinates
    st = run(env, body_for("S", pv["guess"], skip_existing=False))
    assert st["counts"] == {g.PRECISE: 2, g.NOT_FOUND: 1}
    ws = download(env, env.tmp / "o2.xlsx")["S"]
    assert ws["H2"].value != 18.1
    assert (ws["H4"].value, ws["I4"].value) == (18.2, 73.2)
    assert "Previous coordinates left unchanged" in ws.cell(4, col(ws, "Geocoding_Note")).value


def test_6_unfound_and_review(env):
    rows = [[1, "A", "", "nowhere land", "Pune", "MH", 411001],
            [2, "", "", "", "Pune", "", None],                          # city only -> locality centre
            [3, "C", "", "ambiguous plaza", "Pune", "MH", 411001],
            [4, "D", "", "7 Fort Lane wrongpin", "Pune", "MH", 411001],   # PIN mismatch
            [5, "E", "", "Long street 9", "Pune", "MH", 411001]]
    src = make_xlsx(env.tmp / "c6.xlsx", {"S": (HEAD, rows)})
    upload(env, src)
    st = run(env, body_for("S", STD_MAP))
    by = {r["row"]: r for r in st["results"]}
    assert by[2]["status"] == g.NOT_FOUND and by[2]["lat"] is None
    assert by[3]["status"] == g.REVIEW and by[3]["loc_type"] == "APPROXIMATE"
    assert by[4]["status"] == g.REVIEW and "candidates" in by[4]["note"]
    assert by[5]["status"] == g.REVIEW and "PIN mismatch" in by[5]["note"]
    assert by[6]["status"] == g.APPROXIMATE


def test_7_invalid_key(env, monkeypatch):
    monkeypatch.setitem(appmod.KEY, "value", "bad-key")
    src = make_xlsx(env.tmp / "c7.xlsx", {"S": (HEAD, ROWS)})
    upload(env, src)
    st = run(env, body_for("S", STD_MAP))
    assert st["state"] == "error" and "REQUEST_DENIED" in st["error"]
    assert "bad-key" not in st["error"]
    assert st["counts"] == {} and len(fake_google.calls) == 1  # stopped immediately


def test_7b_missing_key(env, monkeypatch):
    monkeypatch.setitem(appmod.KEY, "value", "")
    upload(env, make_xlsx(env.tmp / "c.xlsx", {"S": (HEAD, ROWS)}))
    r = env.post("/api/start", json=body_for("S", STD_MAP))
    assert r.status_code == 400 and "API key" in r.json["error"]


def test_key_entry_in_page(env, monkeypatch):
    srv = fake_google.serve(8998)
    monkeypatch.setenv("GEOCODING_API_URL", "http://127.0.0.1:8998/")
    monkeypatch.setitem(appmod.KEY, "value", "")
    try:
        assert env.get("/api/config").json == {"key_configured": False}
        r = env.post("/api/key", json={"key": "AIza-bad-key-0000000000000"})
        assert r.status_code == 400 and "REQUEST_DENIED" in r.json["error"] and "AIza" not in r.json["error"]
        assert env.post("/api/key", json={"key": "short"}).status_code == 400
        r = env.post("/api/key", json={"key": "AIzaGoodTestKey1234567890"})
        assert r.json == {"key_configured": True}  # key itself never returned
        assert appmod.KEY["value"] == "AIzaGoodTestKey1234567890"
        assert "AIza" not in env.get("/api/config").get_data(as_text=True)
    finally:
        srv.shutdown()


def test_8_quota_then_resume(env):
    rows = [ROWS[0], [2, "Q", "", "quota road 1", "Pune", "MH", 411001], ROWS[1]]
    src = make_xlsx(env.tmp / "c8.xlsx", {"S": (HEAD, rows)})
    upload(env, src)
    body = body_for("S", STD_MAP)
    st = run(env, body)
    assert st["state"] == "error" and "OVER_QUERY_LIMIT" in st["error"]
    assert st["counts"] == {g.PRECISE: 1}  # row 2 kept; row 3 not reached
    n = len(fake_google.calls)
    assert n == 1 + 4  # 1 ok + (1 + 3 retries)
    # "quota lifted": resume without repeating the successful request
    fake_google.calls.clear()
    orig = fake_google.respond

    def lifted(params):
        params = dict(params, address=params["address"].replace("quota", "q"))
        return orig(params)
    env.holder["geo"].session = type("S", (), {"get": lambda self, url, params=None, timeout=None:
                                                fake_google.FakeResponse(lifted(params))})()
    st = run(env, body)
    assert st["state"] == "done" and st["counts"] == {g.PRECISE: 3}
    assert len(fake_google.calls) == 2  # only rows 3 and 4


def test_9_multiple_worksheets(env):
    src = make_xlsx(env.tmp / "c9.xlsx", {"Summary": (["x", "y"], [[1, 2]]), "Customers": (HEAD, ROWS[:1])})
    upload(env, src)
    preview(env, "Customers")
    st = run(env, body_for("Customers", STD_MAP))
    assert st["counts"] == {g.PRECISE: 1}
    wb = download(env, env.tmp / "o.xlsx")
    assert wb.sheetnames == ["Summary", "Customers"]
    assert wb["Summary"].max_column == 2 and wb["Customers"].max_column == 7 + 7


def test_10_large_file(env):
    rows = [[i, f"Outlet {i}", f"98{i:08d}", f"Shop {i}, Lane {i % 50}", "Pune", "Maharashtra", 411001]
            for i in range(3000)]
    src = make_xlsx(env.tmp / "big.xlsx", {"S": (HEAD, rows)})
    upload(env, src)
    st = run(env, body_for("S", STD_MAP))
    assert st["state"] == "done" and st["counts"] == {g.PRECISE: 3000}
    ws = download(env, env.tmp / "o.xlsx")["S"]
    lat = col(ws, "Latitude")
    for i in (0, 1234, 2999):
        q = f"Shop {i}, Lane {i % 50}, Pune, Maharashtra, 411001, India".lower()
        assert ws.cell(i + 2, lat).value == pytest.approx(fake_google._coord(q)[0], abs=1e-6)
        assert ws.cell(i + 2, 3).value == f"98{i:08d}"


def test_retry_failed_rows_bypasses_cache(env):
    src = make_xlsx(env.tmp / "r.xlsx", {"S": (HEAD, ROWS)})
    upload(env, src)
    body = body_for("S", STD_MAP)
    run(env, body)
    fake_google.calls.clear()
    st = run(env, body, mode="retry")
    assert fake_google.calls == ["nowhere village xyz, Pune, Maharashtra, 411001, India",
                                 "Lost Dhaba, nowhere village xyz, Pune, Maharashtra, 411001, India"]
    assert st["counts"] == {g.PRECISE: 2, g.NOT_FOUND: 1}


def test_stop_and_clear(env):
    rows = [[i, f"O{i}", "", f"Shop {i} Lane", "Pune", "MH", 411001] for i in range(200)]
    upload(env, make_xlsx(env.tmp / "s.xlsx", {"S": (HEAD, rows)}))
    env.post("/api/start", json=body_for("S", STD_MAP))
    env.post("/api/stop")
    appmod.S.thread.join()
    assert env.get("/api/status").json["state"] in ("stopped", "done")
    env.post("/api/clear")
    assert not os.listdir(appmod.WORK_DIR)
    assert env.get("/api/download").status_code == 400


def test_xls_upload(env):
    xlwt = pytest.importorskip("xlwt")
    book = xlwt.Workbook()
    sh = book.add_sheet("Old")
    for c, h in enumerate(HEAD):
        sh.write(0, c, h)
    for r, row in enumerate(ROWS[:1], start=1):
        for c, v in enumerate(row):
            sh.write(r, c, v)
    p = env.tmp / "old.xls"
    book.save(str(p))
    assert upload(env, p)["xls"] is True
    preview(env, "Old")
    st = run(env, body_for("Old", STD_MAP))
    assert st["counts"] == {g.PRECISE: 1}
    ws = download(env, env.tmp / "o.xlsx")["Old"]
    assert ws["C2"].value == "9876500001" and ws["G2"].value == 411001


def test_bad_upload(env):
    r = env.post("/api/upload", data={"file": (io.BytesIO(b"not excel"), "x.xlsx")})
    assert r.status_code == 400 and "Traceback" not in r.json["error"]
    r = env.post("/api/upload", data={"file": (io.BytesIO(b"a,b"), "x.csv")})
    assert r.status_code == 400
